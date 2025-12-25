import io
import re
import zipfile
import pathlib
from pathlib import Path
from collections import Counter

# --- 核心依赖导入 ---
try:
    from PyPDF2 import PdfReader
except ImportError:
    PdfReader = None

try:
    from docx import Document
except ImportError:
    Document = None

try:
    from pptx import Presentation
except ImportError:
    Presentation = None

# --- OCR 依赖导入 ---
try:
    import pytesseract
    from pdf2image import convert_from_path
    from PIL import Image, ImageFilter, ImageOps

    OCR_AVAILABLE = True
except ImportError:
    OCR_AVAILABLE = False


class ParserError(Exception):
    pass


class Parser:
    """返回结构化解析结果的解析器（支持 PDF/Word/PPT/MD/图片 及 OCR 识别）。"""

    IMAGE_EXTS = {'.png', '.jpg', '.jpeg', '.tiff', '.tif', '.bmp', '.webp'}

    def __init__(self, poppler_path: str = None, tesseract_cmd: str = None):
        """
        初始化解析器。
        :param poppler_path: Windows 用户需指定 Poppler 的 bin 目录路径。
        :param tesseract_cmd: Windows 用户若未将 Tesseract 加入环境变量，需指定 tesseract.exe 的完整路径。
        """
        self.poppler_path = poppler_path

        if tesseract_cmd and OCR_AVAILABLE:
            pytesseract.pytesseract.tesseract_cmd = tesseract_cmd

    def parse_file(self, file_path, ocr: bool = True, ocr_lang: str = "chi_sim+eng"):
        """
        解析文件入口。
        """
        file_path = Path(file_path)
        if not file_path.exists():
            raise FileNotFoundError(f"{file_path} 不存在")

        meta = {'source': str(file_path)}
        ext = file_path.suffix.lower()
        meta['ext'] = ext
        pages = []

        # 检查 OCR 依赖状态
        actual_ocr = ocr and OCR_AVAILABLE
        if ocr and not OCR_AVAILABLE:
            print("Warning: OCR 已启用但缺少依赖库 (pytesseract/PIL/pdf2image)，将跳过 OCR。")

        if ext == '.pdf':
            if PdfReader is None:
                raise ParserError('缺少 PyPDF2 库。')
            pages = self._parse_pdf(file_path, ocr=actual_ocr, ocr_lang=ocr_lang)

        elif ext == '.docx':
            if Document is None:
                raise ParserError('缺少 python-docx 库。')
            text_content = self._parse_word_text(file_path)
            image_content = self._parse_docx_images(file_path, actual_ocr, ocr_lang) if actual_ocr else ""
            combined = (text_content or "") + ("\n" + image_content if image_content else "")
            pages = [combined]

        elif ext == '.pptx':
            if Presentation is None:
                raise ParserError('缺少 python-pptx 库。')
            pages = self._parse_ppt(file_path, ocr=actual_ocr, ocr_lang=ocr_lang)

        elif ext in ('.md', '.markdown'):
            pages = [self._parse_markdown(file_path, ocr=actual_ocr, ocr_lang=ocr_lang)]

        elif ext in self.IMAGE_EXTS:
            pages = [self._parse_image(file_path, ocr=actual_ocr, ocr_lang=ocr_lang)]

        elif ext in ('.doc', '.ppt'):
            raise ParserError("不支持旧版二进制 .doc / .ppt 格式，请先转换为 .docx / .pptx。")
        else:
            raise ParserError(f"不支持的文件格式: {ext}")

        cleaned_pages = self._clean_pages(pages)
        text = "\n\n".join(cleaned_pages)

        return {
            'text': text,
            'pages': cleaned_pages,
            'meta': meta,
        }

    # ========================== OCR 核心辅助 ==========================

    def _preprocess_image_for_ocr(self, img):
        """图像预处理：转灰度、增强对比度。"""
        try:
            if img.mode != 'RGB':
                img = img.convert('RGB')
            w, h = img.size
            if max(w, h) < 1000:
                img = img.resize((w * 2, h * 2), Image.LANCZOS)
            img = ImageOps.grayscale(img)
            img = img.filter(ImageFilter.MedianFilter(size=3))
            img = ImageOps.autocontrast(img)
            return img
        except Exception:
            return img

    def _ocr_image_bytes(self, image_data, lang: str = "chi_sim+eng"):
        """处理二进制图片流。"""
        if not OCR_AVAILABLE or not image_data:
            return ""
        try:
            with io.BytesIO(image_data) as f:
                img = Image.open(f)
                img = self._preprocess_image_for_ocr(img)
                text = pytesseract.image_to_string(img, lang=lang)
                return text.strip()
        except Exception:
            return ""

    # ========================== 格式解析实现 ==========================

    def _parse_pdf(self, file_path: Path, ocr: bool, ocr_lang: str):
        reader = PdfReader(str(file_path))
        pages_text = []

        for i, page in enumerate(reader.pages, start=1):
            parts = []
            # 1. 尝试提取常规文本
            try:
                text = page.extract_text() or ""
                if text.strip():
                    parts.append(text)
            except Exception:
                pass

            # 2. 提取并识别页面内的嵌入图片
            if ocr:
                try:
                    if hasattr(page, 'images') and page.images:
                        for img_obj in page.images:
                            extracted = self._ocr_image_bytes(img_obj.data, lang=ocr_lang)
                            if extracted:
                                parts.append(f"[Image Text]: {extracted}")
                except Exception:
                    pass

            # 3. 兜底策略：如果页面无文字，则将整页渲染成图片进行 OCR
            full_page_text = "\n".join(parts)
            if ocr and len(full_page_text.strip()) < 10:
                try:
                    # 关键修复点：这里传入了 self.poppler_path
                    pil_pages = convert_from_path(
                        str(file_path),
                        first_page=i,
                        last_page=i,
                        poppler_path=self.poppler_path
                    )
                    if pil_pages:
                        img = self._preprocess_image_for_ocr(pil_pages[0])
                        fallback_text = pytesseract.image_to_string(img, lang=ocr_lang)
                        full_page_text = fallback_text.strip()
                except Exception as e:
                    print(f"Warning: PDF Page {i} OCR failed. Error: {e}")

            header = f"---PAGE {i}---\n"
            pages_text.append(header + full_page_text.strip())

        return pages_text

    def _parse_word_text(self, file_path: Path) -> str:
        doc = Document(str(file_path))
        return "\n".join([p.text for p in doc.paragraphs])

    def _parse_docx_images(self, file_path: Path, ocr: bool, ocr_lang: str) -> str:
        extracted_texts = []
        try:
            with zipfile.ZipFile(file_path) as z:
                media_files = [f for f in z.namelist() if f.startswith('word/media/')]
                for media in media_files:
                    if any(media.lower().endswith(ext) for ext in ['.png', '.jpg', '.jpeg', '.bmp']):
                        img_data = z.read(media)
                        text = self._ocr_image_bytes(img_data, lang=ocr_lang)
                        if text:
                            extracted_texts.append(f"---Embedded Image ({Path(media).name})---\n{text}")
        except Exception:
            pass
        return "\n".join(extracted_texts)

    def _parse_ppt(self, file_path: Path, ocr: bool, ocr_lang: str):
        prs = Presentation(str(file_path))
        slides = []
        for si, slide in enumerate(prs.slides, start=1):
            texts = []
            for shape in slide.shapes:
                if hasattr(shape, 'text') and shape.text.strip():
                    texts.append(shape.text.strip())
                if hasattr(shape, 'has_table') and shape.has_table:
                    for r in shape.table.rows:
                        texts.append("\t".join(cell.text.strip() for cell in r.cells))
                if ocr and hasattr(shape, 'image'):
                    img_text = self._ocr_image_bytes(shape.image.blob, lang=ocr_lang)
                    if img_text:
                        texts.append(f"[Slide Image]: {img_text}")
            slides.append(f"---SLIDE {si}---\n" + "\n".join(texts))
        return slides

    def _parse_markdown(self, file_path: Path, ocr: bool, ocr_lang: str) -> str:
        content = ""
        for enc in ("utf-8", "utf-8-sig", "gbk"):
            try:
                content = file_path.read_text(encoding=enc)
                break
            except UnicodeDecodeError:
                continue

        if ocr and content:
            base_dir = file_path.parent
            img_pattern = re.compile(r'!\[.*?\]\((.*?)\)')

            def replace_with_ocr(match):
                rel_path = match.group(1).split()[0]
                img_path = (base_dir / rel_path).resolve()
                if img_path.exists() and img_path.suffix.lower() in self.IMAGE_EXTS:
                    text = self._ocr_image_bytes(img_path.read_bytes(), lang=ocr_lang)
                    if text: return f"{match.group(0)}\n[Image OCR]: {text}\n"
                return match.group(0)

            content = img_pattern.sub(replace_with_ocr, content)
        return content

    def _parse_image(self, file_path: Path, ocr: bool, ocr_lang: str) -> str:
        if not ocr:
            return f"---IMAGE {file_path.name} (OCR Disabled)---"
        try:
            img = Image.open(str(file_path))
            proc = self._preprocess_image_for_ocr(img)
            text = pytesseract.image_to_string(proc, lang=ocr_lang)
            return f"---IMAGE {file_path.name}---\n{text.strip()}"
        except Exception as e:
            return f"Image Parse Error: {e}"

    def _clean_pages(self, pages):
        """简单的页眉页脚过滤逻辑。"""
        if not pages: return []
        pages_lines = [[ln.strip() for ln in p.splitlines() if ln.strip()] for p in pages]
        line_page_count = Counter()
        for lines in pages_lines:
            for ln in set(lines): line_page_count[ln] += 1

        num_pages = len(pages)
        remove_lines = {ln for ln, cnt in line_page_count.items() if num_pages > 2 and (cnt / num_pages) > 0.6}

        cleaned = []
        for lines in pages_lines:
            filtered = [ln for ln in lines if ln not in remove_lines]
            cleaned.append("\n".join(filtered if filtered else lines))
        return cleaned


# ========================== 使用示例 ==========================
if __name__ == "__main__":
    # 配置路径 (根据你的实际安装位置修改)
    MY_POPPLER_PATH = r'C:\Program Files\poppler\Library\bin'
    MY_TESSERACT_EXE = r'C:\Program Files\Tesseract-OCR\tesseract.exe'

    # 初始化解析器
    parser = Parser(poppler_path=MY_POPPLER_PATH, tesseract_cmd=MY_TESSERACT_EXE)

    try:
        # 测试解析
        result = parser.parse_file("your_file.pdf", ocr=True)
        print("--- 解析成功 ---")
        print(result['text'][:1000])  # 打印前1000字
    except Exception as e:
        print(f"解析失败: {e}")
